#!/usr/bin/env python3
"""Render presentation.html slide by slide and package it as PDF and PPTX.

Requires Google Chrome plus `pip install python-pptx pillow`.
Run from anywhere:  python3 tools/build_downloads.py
Re-run after editing presentation.html so the downloads stay in sync.
"""
import functools
import http.server
import io
import os
import re
import subprocess
import sys
import tempfile
import threading
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "downloads"
NAME = "carp-disease-presentation"
TITLE = "Investigation and Molecular Characterization of Diseases Affecting Cultured Carps in Bangladesh"
W, H = 1920, 1080
CHROME = os.environ.get("CHROME", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")


class QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def serve():
    handler = functools.partial(QuietHandler, directory=str(ROOT))
    httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd


def capture(port, n, path):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=1",
                    f"--window-size={W},{H}", "--virtual-time-budget=5000", f"--screenshot={path}",
                    f"http://127.0.0.1:{port}/presentation.html?capture#{n}"],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def main():
    count = len(re.findall(r'<section class="slide', (ROOT / "presentation.html").read_text()))
    httpd = serve()
    port = httpd.server_address[1]
    OUT.mkdir(exist_ok=True)
    images = []
    with tempfile.TemporaryDirectory() as tmp:
        for n in range(1, count + 1):
            png = Path(tmp) / f"s{n}.png"
            capture(port, n, png)
            images.append(Image.open(png).convert("RGB").crop((0, 0, W, H)))
            print(f"  slide {n}/{count}")
    httpd.shutdown()

    # PDF: one 16:9 page per slide
    pdf = OUT / f"{NAME}.pdf"
    images[0].save(pdf, "PDF", save_all=True, append_images=images[1:], resolution=144, quality=88, title=TITLE)

    # PPTX: 16:9 deck, each slide a full-bleed picture
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    prs.core_properties.title = TITLE
    blank = prs.slide_layouts[6]
    for img in images:
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=88)
        buf.seek(0)
        prs.slides.add_slide(blank).shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
    pptx = OUT / f"{NAME}.pptx"
    prs.save(pptx)

    for f in (pdf, pptx):
        print(f"{f.relative_to(ROOT)}  {f.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    sys.exit(main())
